#!/usr/bin/env -S uv run
# /// script
# dependencies = ["openpyxl>=3.1.5", "PyYAML>=6.0"]
# ///

from __future__ import annotations

import os
import re
import sqlite3
import time
import unicodedata
import warnings
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import ClassVar
from urllib.parse import urlparse

from openpyxl import load_workbook
from shared import ROOT
from shared.data_loader import ManualDataLoader
from shared.http import JsonSource

warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
    module="openpyxl.worksheet._reader",
)

FISKE_PATH = ROOT / "Fiske 2024 Summary.xlsx"
CIP_PATH = ROOT / "CIP codes bio.xlsx"
CIP_TO_SUBFIELD_PATH = ROOT / "cip_to_openalex_subfields_openalex_native.md"
SCORECARD_API_URL = "https://api.data.gov/ed/collegescorecard/v1/schools"
IPEDS_INSTITUTION_PROFILE_URL = "https://nces.ed.gov/ipeds/institution-profile/{unitid}"
ROR_API_URL = "https://api.ror.org/v2/organizations"
OPENALEX_INSTITUTIONS_API_URL = "https://api.openalex.org/institutions"
OPENALEX_FIELDS_API_URL = "https://api.openalex.org/fields"
NAME_STOPWORDS = {"a", "an", "and", "at", "for", "in", "of", "the"}
US_STATE_CODES = {
    "AL",
    "AK",
    "AZ",
    "AR",
    "CA",
    "CO",
    "CT",
    "DC",
    "DE",
    "FL",
    "GA",
    "HI",
    "IA",
    "ID",
    "IL",
    "IN",
    "KS",
    "KY",
    "LA",
    "MA",
    "MD",
    "ME",
    "MI",
    "MN",
    "MO",
    "MS",
    "MT",
    "NC",
    "ND",
    "NE",
    "NH",
    "NJ",
    "NM",
    "NV",
    "NY",
    "OH",
    "OK",
    "OR",
    "PA",
    "PR",
    "RI",
    "SC",
    "SD",
    "TN",
    "TX",
    "UT",
    "VA",
    "VT",
    "WA",
    "WI",
    "WV",
    "WY",
}
MANUAL_HOMEPAGES = {
    "Deep Springs College": {
        "canonical_name": "Deep Springs College",
        "homepage_url": "https://www.deepsprings.edu/",
    }
}

SUPPLEMENTAL_CIPS = [
    ("27.0502", "Mathematical Statistics and Probability"),
    ("27.0503", "Mathematics and Statistics"),
    ("27.0599", "Statistics, Other"),
    ("27.0601", "Applied Statistics, General"),
]


@dataclass(frozen=True)
class UniversitySourceRow:
    fiske_name: str
    state_code: str
    ipeds_unitid: int | None


@dataclass(frozen=True)
class UniversityRecord:
    canonical_name: str
    country_code: str
    ipeds_unitid: int | None
    homepage_url: str


@dataclass(frozen=True)
class CipRecord:
    cip_code: str
    cip_title: str


@dataclass(frozen=True)
class IdentifierRecord:
    system_code: str
    external_id: str
    external_url: str | None


@dataclass(frozen=True)
class OaFieldRecord:
    field_id: int
    domain_id: int
    field_name: str


@dataclass(frozen=True)
class OaSubfieldRecord:
    subfield_id: int
    field_id: int
    subfield_name: str


@dataclass(frozen=True)
class CipTrainsSubfieldRecord:
    cip_code: str
    subfield_id: int
    fit_level: str


def chunked(values: list[str], size: int) -> list[list[str]]:
    return [values[index : index + size] for index in range(0, len(values), size)]


def with_openalex_api_key(params: dict[str, str | int]) -> dict[str, str | int]:
    openalex_api_key = os.environ.get("OPENALEX_API_KEY")
    if not openalex_api_key:
        return params
    return {**params, "api_key": openalex_api_key}


def normalize_host(value: str) -> str:
    host = value.lower().strip()
    if host.startswith("www."):
        host = host[4:]
    return host


def normalize_homepage_url(url: str) -> str:
    cleaned = url.strip()
    if not cleaned:
        raise ValueError("homepage_url cannot be blank")
    if not urlparse(cleaned).scheme:
        return f"https://{cleaned.lstrip('/')}"
    return cleaned


def homepage_host(url: str) -> str:
    return normalize_host(urlparse(normalize_homepage_url(url)).netloc.split(":", 1)[0])


def host_matches(left: str, right: str) -> bool:
    lhs = normalize_host(left)
    rhs = normalize_host(right)
    return lhs == rhs or lhs.endswith(f".{rhs}") or rhs.endswith(f".{lhs}")


def host_match_score(left: str, right: str) -> int:
    lhs = normalize_host(left)
    rhs = normalize_host(right)
    if lhs == rhs:
        return 100
    if lhs.endswith(f".{rhs}") or rhs.endswith(f".{lhs}"):
        return 60
    return 0


def canonicalize_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    for dash in ("\u2013", "\u2014", "\u2212", "\u2011"):
        normalized = normalized.replace(dash, "-")
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    lower_value = ascii_value.lower().replace("&", " and ")
    collapsed = re.sub(r"[^a-z0-9]+", " ", lower_value).strip()
    if collapsed.startswith("the "):
        collapsed = collapsed[4:]
    return re.sub(r"\s+", " ", collapsed)


def canonicalize_name_without_stopwords(value: str) -> str:
    return " ".join(
        token for token in canonicalize_name(value).split() if token not in NAME_STOPWORDS
    )


def build_identifier_external_url(system_code: str, external_id: str) -> str | None:
    if system_code == "IPEDS":
        return IPEDS_INSTITUTION_PROFILE_URL.format(unitid=external_id)
    if system_code == "ROR":
        return f"https://ror.org/{external_id}"
    if system_code == "OPENALEX":
        return f"https://openalex.org/{external_id}"
    return None


def load_us_universities(path: Path) -> list[UniversitySourceRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active

    universities: list[UniversitySourceRow] = []
    seen: set[tuple[str, str] | int] = set()

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        fiske_name = str(row[1]).strip() if row[1] else ""
        state_code = str(row[2]).strip() if row[2] else ""
        if not fiske_name or state_code not in US_STATE_CODES:
            continue

        ipeds_unitid = int(row[16]) if row[16] is not None else None
        dedupe_key: tuple[str, str] | int = (
            (fiske_name, state_code) if ipeds_unitid is None else ipeds_unitid
        )

        if dedupe_key in seen:
            continue

        seen.add(dedupe_key)
        universities.append(
            UniversitySourceRow(
                fiske_name=fiske_name,
                state_code=state_code,
                ipeds_unitid=ipeds_unitid,
            )
        )

    workbook.close()
    return universities


def fetch_scorecard_records(
    source: JsonSource, unitids: list[int]
) -> dict[int, dict[str, str | int]]:
    api_key = os.environ.get("SCORECARD_API_KEY")
    if not api_key:
        raise RuntimeError(
            "SCORECARD_API_KEY is required to enrich universities with homepage URLs."
        )

    by_unitid: dict[int, dict[str, str | int]] = {}
    for id_chunk in chunked([str(unitid) for unitid in unitids], size=50):
        payload = source.fetch(
            SCORECARD_API_URL,
            params={
                "api_key": api_key,
                "fields": "id,school.name,school.state,school.school_url",
                "id": ",".join(id_chunk),
                "per_page": len(id_chunk),
            },
        )

        for result in payload["results"]:
            by_unitid[int(result["id"])] = result

    missing_ids = sorted(set(unitids) - set(by_unitid))
    if missing_ids:
        raise RuntimeError(f"College Scorecard did not return data for unitids: {missing_ids}")

    missing_urls = [
        unitid
        for unitid, result in by_unitid.items()
        if not str(result.get("school.school_url", "")).strip()
    ]
    if missing_urls:
        raise RuntimeError(
            f"College Scorecard returned blank school.school_url for unitids: {missing_urls}"
        )

    return by_unitid


def enrich_universities(
    source: JsonSource, rows: list[UniversitySourceRow]
) -> list[UniversityRecord]:
    scorecard_rows = fetch_scorecard_records(
        source, [row.ipeds_unitid for row in rows if row.ipeds_unitid is not None]
    )

    universities: list[UniversityRecord] = []
    for row in rows:
        if row.ipeds_unitid is not None:
            scorecard = scorecard_rows[row.ipeds_unitid]
            canonical_name = str(scorecard["school.name"]).strip()
            scorecard_state = str(scorecard["school.state"]).strip()
            if scorecard_state != row.state_code:
                raise RuntimeError(
                    f"State mismatch for unitid {row.ipeds_unitid}: "
                    f"Fiske={row.state_code}, Scorecard={scorecard_state}"
                )
            homepage_url = normalize_homepage_url(str(scorecard["school.school_url"]))
        else:
            manual = MANUAL_HOMEPAGES.get(row.fiske_name)
            if manual is None:
                raise RuntimeError(
                    f"No manual homepage mapping is defined for {row.fiske_name} ({row.state_code})."
                )
            canonical_name = manual["canonical_name"]
            homepage_url = normalize_homepage_url(manual["homepage_url"])

        universities.append(
            UniversityRecord(
                canonical_name=canonical_name,
                country_code="US",
                ipeds_unitid=row.ipeds_unitid,
                homepage_url=homepage_url,
            )
        )

    universities.sort(key=lambda item: item.canonical_name)
    return universities


def ror_hosts(record: dict) -> set[str]:
    hosts = {normalize_host(domain) for domain in record.get("domains", []) if domain}
    for link in record.get("links", []):
        if link.get("type") == "website" and link.get("value"):
            hosts.add(homepage_host(link["value"]))
    return hosts


def score_ror_candidate(university: UniversityRecord, record: dict) -> tuple[int, bool, bool]:
    score = 0
    normalized_name = canonicalize_name(university.canonical_name)
    normalized_name_no_stopwords = canonicalize_name_without_stopwords(university.canonical_name)
    candidate_names = [
        canonicalize_name(name["value"]) for name in record.get("names", []) if name.get("value")
    ]
    candidate_names_no_stopwords = [
        canonicalize_name_without_stopwords(name["value"])
        for name in record.get("names", [])
        if name.get("value")
    ]
    exact_name_match = normalized_name in candidate_names
    if exact_name_match:
        score += 80
    elif normalized_name_no_stopwords in candidate_names_no_stopwords:
        score += 70
    elif any(
        normalized_name in candidate_name or candidate_name in normalized_name
        for candidate_name in candidate_names
    ):
        score += 20

    domain_score = max(
        (
            host_match_score(homepage_host(university.homepage_url), candidate_host)
            for candidate_host in ror_hosts(record)
        ),
        default=0,
    )
    domain_match = domain_score > 0
    score += domain_score

    if "education" in record.get("types", []):
        score += 15
    if record.get("status") == "active":
        score += 5
    locations = record.get("locations", [])
    if any(
        location.get("geonames_details", {}).get("country_code") == "US" for location in locations
    ):
        score += 10
    else:
        score -= 50

    return score, exact_name_match, domain_match


def fetch_ror_record(source: JsonSource, university: UniversityRecord) -> dict:
    candidates: dict[str, dict] = {}
    for query in (university.canonical_name, homepage_host(university.homepage_url)):
        if not query:
            continue
        payload = source.fetch(ROR_API_URL, params={"query": query})
        for item in payload.get("items", []):
            candidates[item["id"]] = item
        time.sleep(0.05)

    if not candidates:
        raise RuntimeError(f"No ROR candidates found for {university.canonical_name}")

    ranked: list[tuple[int, bool, bool, dict]] = []
    for candidate in candidates.values():
        score, exact_name_match, domain_match = score_ror_candidate(university, candidate)
        ranked.append((score, exact_name_match, domain_match, candidate))
    ranked.sort(
        key=lambda item: (
            item[0],
            item[1],
            item[2],
            item[3]["id"],
        ),
        reverse=True,
    )

    top_score, top_exact_name, top_domain, top_record = ranked[0]
    second_score = ranked[1][0] if len(ranked) > 1 else -(10**9)
    if top_score < 110 or (top_score - second_score < 20 and not (top_exact_name and top_domain)):
        raise RuntimeError(
            f"Ambiguous ROR match for {university.canonical_name}: "
            f"top={top_record['id']} score={top_score}, second={second_score}"
        )
    return top_record


def fetch_openalex_record(source: JsonSource, university: UniversityRecord, ror_url: str) -> dict:
    params: dict[str, str | int] = {
        "filter": f"ror:{ror_url}",
        "select": "id,display_name,ror,country_code",
    }
    payload = source.fetch(OPENALEX_INSTITUTIONS_API_URL, params=with_openalex_api_key(params))
    results = payload.get("results", [])
    if len(results) == 1:
        return results[0]
    if len(results) > 1:
        raise RuntimeError(
            f"Multiple OpenAlex institutions found for ROR {ror_url}: "
            f"{[result['id'] for result in results]}"
        )

    fallback_params: dict[str, str | int] = {
        "search": university.canonical_name,
        "select": "id,display_name,ror,country_code",
    }
    fallback = source.fetch(
        OPENALEX_INSTITUTIONS_API_URL,
        params=with_openalex_api_key(fallback_params),
    )
    fallback_results = [
        result for result in fallback.get("results", []) if result.get("ror") == ror_url
    ]
    if len(fallback_results) == 1:
        return fallback_results[0]

    raise RuntimeError(
        f"No OpenAlex institution found for {university.canonical_name} using ROR {ror_url}"
    )


def build_identifier_records(
    source: JsonSource, university: UniversityRecord
) -> list[IdentifierRecord]:
    records: list[IdentifierRecord] = []
    if university.ipeds_unitid is not None:
        external_id = str(university.ipeds_unitid)
        records.append(
            IdentifierRecord(
                system_code="IPEDS",
                external_id=external_id,
                external_url=build_identifier_external_url("IPEDS", external_id),
            )
        )

    ror_record = fetch_ror_record(source, university)
    ror_url = ror_record["id"]
    ror_id = ror_url.rsplit("/", 1)[1]
    records.append(
        IdentifierRecord(
            system_code="ROR",
            external_id=ror_id,
            external_url=build_identifier_external_url("ROR", ror_id),
        )
    )

    openalex_record = fetch_openalex_record(source, university, ror_url)
    openalex_id = openalex_record["id"].rsplit("/", 1)[1]
    records.append(
        IdentifierRecord(
            system_code="OPENALEX",
            external_id=openalex_id,
            external_url=build_identifier_external_url("OPENALEX", openalex_id),
        )
    )

    return records


def load_openalex_hierarchy(
    source: JsonSource,
) -> tuple[list[OaFieldRecord], list[OaSubfieldRecord]]:
    oa_fields: list[OaFieldRecord] = []
    oa_subfields: list[OaSubfieldRecord] = []
    seen_fields: set[int] = set()
    seen_subfields: set[int] = set()

    cursor = "*"
    while cursor:
        payload = source.fetch(
            OPENALEX_FIELDS_API_URL,
            params=with_openalex_api_key({"per-page": 200, "cursor": cursor}),
        )
        for field in payload.get("results", []):
            field_id = int(str(field["id"]).rsplit("/", 1)[1])
            if field_id not in seen_fields:
                seen_fields.add(field_id)
                domain_id = int(str(field["domain"]["id"]).rsplit("/", 1)[1])
                oa_fields.append(
                    OaFieldRecord(
                        field_id=field_id,
                        domain_id=domain_id,
                        field_name=str(field["display_name"]).strip(),
                    )
                )

            for subfield in field.get("subfields", []):
                subfield_id = int(str(subfield["id"]).rsplit("/", 1)[1])
                if subfield_id in seen_subfields:
                    continue
                seen_subfields.add(subfield_id)
                oa_subfields.append(
                    OaSubfieldRecord(
                        subfield_id=subfield_id,
                        field_id=field_id,
                        subfield_name=str(subfield["display_name"]).strip(),
                    )
                )
        cursor = payload.get("meta", {}).get("next_cursor")

    oa_fields.sort(key=lambda item: item.field_id)
    oa_subfields.sort(key=lambda item: item.subfield_id)
    return oa_fields, oa_subfields


def format_cip_code(value: object) -> str:
    if value is None:
        raise ValueError("CIP code cannot be null")
    if isinstance(value, str):
        text = value.strip()
        if not text:
            raise ValueError("CIP code cannot be blank")
        return text
    return format(Decimal(str(value)).quantize(Decimal("0.0000")), "f")


def load_cips(path: Path) -> list[CipRecord]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    worksheet = workbook.active

    cips: list[CipRecord] = []
    seen_codes: set[str] = set()

    for code_cell, title_cell in worksheet.iter_rows(min_row=2, min_col=4, max_col=5):
        if code_cell.value is None and title_cell.value is None:
            continue

        cip_code = format_cip_code(code_cell.value)
        cip_title = str(title_cell.value).strip() if title_cell.value else ""
        if not cip_title:
            raise ValueError(f"Missing CIP title for code {cip_code}")
        if cip_code in seen_codes:
            raise ValueError(f"Duplicate CIP code encountered: {cip_code}")

        seen_codes.add(cip_code)
        cips.append(CipRecord(cip_code=cip_code, cip_title=cip_title))

    for cip_code, cip_title in SUPPLEMENTAL_CIPS:
        if cip_code in seen_codes:
            continue
        seen_codes.add(cip_code)
        cips.append(CipRecord(cip_code=cip_code, cip_title=cip_title))

    workbook.close()
    cips.sort(key=lambda item: item.cip_code)
    return cips


def load_cip_trains_subfield(
    path: Path, known_cip_codes: set[str]
) -> list[CipTrainsSubfieldRecord]:
    records: list[CipTrainsSubfieldRecord] = []
    seen_pairs: set[tuple[str, int]] = set()
    valid_fit_levels = {"primary", "secondary", "adjacent"}
    mapping_header = [
        "CIP code",
        "CIP title",
        "OpenAlex subfield ID",
        "OpenAlex subfield name",
        "OpenAlex field",
        "OpenAlex domain",
        "fit",
        "Why",
    ]
    in_mapping_table = False

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            in_mapping_table = False
            continue
        if not line.startswith("|"):
            in_mapping_table = False
            continue

        columns = [column.strip() for column in line.strip("|").split("|")]
        if columns == mapping_header:
            in_mapping_table = True
            continue
        if columns and all(column and set(column) == {"-"} for column in columns):
            continue
        if not in_mapping_table or len(columns) != len(mapping_header):
            continue

        cip_code = columns[0]
        subfield_url = columns[2].strip("`")
        fit_level = columns[6]
        match = re.search(r"/subfields/(\d+)$", subfield_url)
        if not cip_code or match is None or not fit_level:
            continue
        subfield_id = int(match.group(1))

        if cip_code not in known_cip_codes:
            raise ValueError(f"Unknown CIP code in fit map: {cip_code}")
        if fit_level not in valid_fit_levels:
            raise ValueError(
                f"Unexpected fit_level {fit_level!r} for CIP {cip_code} and subfield {subfield_id}"
            )

        pair = (cip_code, subfield_id)
        if pair in seen_pairs:
            raise ValueError(f"Duplicate CIP/subfield mapping encountered: {pair}")
        seen_pairs.add(pair)
        records.append(
            CipTrainsSubfieldRecord(
                cip_code=cip_code,
                subfield_id=subfield_id,
                fit_level=fit_level,
            )
        )

    records.sort(key=lambda item: (item.cip_code, item.subfield_id))
    return records


def initialize_database(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        PRAGMA foreign_keys = ON;

        DROP TABLE IF EXISTS university_identifiers;
        DROP TABLE IF EXISTS identifier_systems;
        DROP TABLE IF EXISTS oa_subfields;
        DROP TABLE IF EXISTS oa_fields;
        DROP TABLE IF EXISTS cip_trains_subfield;
        DROP TABLE IF EXISTS universities;
        DROP TABLE IF EXISTS cips;

        CREATE TABLE universities (
            university_id INTEGER PRIMARY KEY,
            canonical_name TEXT NOT NULL,
            country_code TEXT NOT NULL CHECK (country_code = 'US'),
            homepage_url TEXT NOT NULL
        );

        CREATE INDEX idx_universities_canonical_name
            ON universities (canonical_name);

        CREATE TABLE identifier_systems (
            system_code TEXT PRIMARY KEY,
            system_name TEXT NOT NULL
        );

        CREATE TABLE university_identifiers (
            university_id INTEGER NOT NULL,
            system_code TEXT NOT NULL,
            external_id TEXT NOT NULL,
            external_url TEXT,
            PRIMARY KEY (university_id, system_code),
            UNIQUE (system_code, external_id),
            FOREIGN KEY (university_id) REFERENCES universities (university_id),
            FOREIGN KEY (system_code) REFERENCES identifier_systems (system_code)
        );

        CREATE TABLE cips (
            cip_code TEXT PRIMARY KEY,
            cip_title TEXT NOT NULL
        );

        CREATE TABLE oa_fields (
            field_id INTEGER PRIMARY KEY,
            domain_id INTEGER NOT NULL,
            field_name TEXT NOT NULL
        );

        CREATE TABLE oa_subfields (
            subfield_id INTEGER PRIMARY KEY,
            field_id INTEGER NOT NULL,
            subfield_name TEXT NOT NULL
        );

        CREATE TABLE cip_trains_subfield (
            cip_code TEXT NOT NULL,
            subfield_id INTEGER NOT NULL,
            fit_level TEXT NOT NULL CHECK (fit_level IN ('primary', 'secondary', 'adjacent')),
            PRIMARY KEY (cip_code, subfield_id),
            FOREIGN KEY (cip_code) REFERENCES cips (cip_code)
        );
        """
    )


def write_database(
    db_path: Path,
    universities: list[UniversityRecord],
    cips: list[CipRecord],
    source: JsonSource,
) -> None:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path)
    try:
        initialize_database(connection)
        oa_fields, oa_subfields = load_openalex_hierarchy(source)
        cts_records = load_cip_trains_subfield(
            CIP_TO_SUBFIELD_PATH,
            {cip.cip_code for cip in cips},
        )
        connection.execute(
            """
            INSERT INTO identifier_systems (system_code, system_name)
            VALUES (?, ?)
            """,
            ("IPEDS", "Integrated Postsecondary Education Data System"),
        )
        connection.executemany(
            """
            INSERT INTO identifier_systems (system_code, system_name)
            VALUES (?, ?)
            """,
            [
                ("ROR", "Research Organization Registry"),
                ("OPENALEX", "OpenAlex"),
            ],
        )

        identifier_rows: list[tuple[int, str, str, str | None]] = []
        for university in universities:
            cursor = connection.execute(
                """
                INSERT INTO universities (
                    canonical_name,
                    country_code,
                    homepage_url
                ) VALUES (?, ?, ?)
                """,
                (
                    university.canonical_name,
                    university.country_code,
                    university.homepage_url,
                ),
            )
            university_id = cursor.lastrowid
            if university_id is None:
                raise RuntimeError("Failed to retrieve lastrowid after university insert")
            for identifier in build_identifier_records(source, university):
                identifier_rows.append(
                    (
                        university_id,
                        identifier.system_code,
                        identifier.external_id,
                        identifier.external_url,
                    )
                )

        connection.executemany(
            """
            INSERT INTO university_identifiers (
                university_id,
                system_code,
                external_id,
                external_url
            ) VALUES (?, ?, ?, ?)
            """,
            identifier_rows,
        )

        connection.executemany(
            """
            INSERT INTO cips (cip_code, cip_title)
            VALUES (?, ?)
            """,
            [(cip.cip_code, cip.cip_title) for cip in cips],
        )

        connection.executemany(
            """
            INSERT INTO cip_trains_subfield (cip_code, subfield_id, fit_level)
            VALUES (?, ?, ?)
            """,
            [(record.cip_code, record.subfield_id, record.fit_level) for record in cts_records],
        )

        connection.executemany(
            """
            INSERT INTO oa_fields (field_id, domain_id, field_name)
            VALUES (?, ?, ?)
            """,
            [(field.field_id, field.domain_id, field.field_name) for field in oa_fields],
        )

        connection.executemany(
            """
            INSERT INTO oa_subfields (subfield_id, field_id, subfield_name)
            VALUES (?, ?, ?)
            """,
            [
                (subfield.subfield_id, subfield.field_id, subfield.subfield_name)
                for subfield in oa_subfields
            ],
        )
        connection.commit()
    finally:
        connection.close()


class BootstrapDbLoader(ManualDataLoader):
    requires_json_source = True
    json_source_headers: ClassVar[dict[str, str]] = {"User-Agent": "univ-bootstrap/1.0"}
    default_request_timeout = 30
    source_code = "BOOTSTRAP_DB"
    description = "Build the local SQLite database for the university project."
    execute_help = "create or refresh the SQLite database (default is dry-run)"
    requires_db = False
    requires_config = False

    def load(
        self,
        *,
        conn,
        config: dict,
        args,
        snapshot: str,
    ) -> tuple[list[dict], dict]:
        universities = enrich_universities(self.json_source, load_us_universities(FISKE_PATH))
        cips = load_cips(CIP_PATH)
        if args.execute:
            write_database(args.db, universities, cips, self.json_source)

        return [], {
            "executed": args.execute,
            "database": str(args.db),
            "universities": len(universities),
            "cips": len(cips),
        }


if __name__ == "__main__":
    raise SystemExit(BootstrapDbLoader().run())
